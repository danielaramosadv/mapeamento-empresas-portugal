"""
Limpeza e padronização da base de empresas (exercício de Análise de Dados e TI aplicada à Gestão).

Uso:
    python limpar_base.py entrada.csv saida.xlsx saida.csv

A base original NÃO está neste repositório. O arquivo de entrada deve ter as colunas:
    name, url, industry, social page, location, founded, size

O que o script faz:
  - remove espaços extras;
  - restaura as faixas de tamanho que o Excel/Sheets converteu em datas ("11-50" -> "nov/50");
  - separa a localização em Cidade, Distrito, Região e País;
  - cria a coluna "Ordem do Tamanho" para ordenar as faixas no painel;
  - gera um .xlsx (com aba de notas) e um .csv limpos.
"""
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

src_csv, out_xlsx, out_csv = sys.argv[1], sys.argv[2], sys.argv[3]

d = pd.read_csv(src_csv, dtype=str, keep_default_na=False)
for c in d.columns:
    d[c] = d[c].str.strip().str.replace(r"\s+", " ", regex=True)

# ---------- Tamanho ----------
# O Excel/Sheets converteu "11-50" em "nov/50" e "1-10" em "01/out." (leu como data).
size_map = {
    "01/out.": ("1 a 10", 1),
    "nov/50": ("11 a 50", 2),
    "51-200": ("51 a 200", 3),
    "201-500": ("201 a 500", 4),
    "501-1000": ("501 a 1.000", 5),
    "1001-5000": ("1.001 a 5.000", 6),
    "5001-10000": ("5.001 a 10.000", 7),
    "10001+": ("Mais de 10.000", 8),
    "Self": ("Autônomo (self-employed)", 0),
}
unknown = set(d["size"]) - set(size_map)
assert not unknown, unknown

# ---------- Localização ----------
district_to_region = {
    "Lisboa": "Área Metropolitana de Lisboa",
    "Setúbal": "Área Metropolitana de Lisboa",
    "Porto": "Norte", "Braga": "Norte", "Viana do Castelo": "Norte", "Vila Real": "Norte",
    "Aveiro": "Centro", "Coimbra": "Centro", "Viseu": "Centro", "Leiria": "Centro",
    "Castelo Branco": "Centro", "Santarém": "Centro",
    "Faro": "Algarve",
    "Açores": "Açores",
}
fix_city = {"Lisbon": "Lisboa"}
fix_district = {"Lisbon": "Lisboa", "Ilha de Santa Maria": "Açores"}


def parse_location(loc):
    """Devolve (cidade, distrito, regiao, pais)."""
    if loc == "Portugal":
        return ("Não informada", "Não informado", "Não informada", "Portugal")
    if loc == "Greater Póvoa de Varzim Area":
        return ("Póvoa de Varzim", "Porto", "Norte", "Portugal")
    if loc == "Porto Metropolitan Area":
        return ("Porto", "Porto", "Norte", "Portugal")
    parts = [p.strip() for p in loc.split(". ")]
    country = parts[-1]
    if country != "Portugal":
        return (parts[0], "N/A (fora de Portugal)", "Fora de Portugal", country)
    city = fix_city.get(parts[0], parts[0])
    if len(parts) == 2:  # "Vila do Porto. Portugal" -> ilha de Santa Maria (Açores)
        dist = "Açores" if parts[0] == "Vila do Porto" else "Não informado"
    else:
        dist = fix_district.get(parts[1], parts[1])
    region = district_to_region.get(dist, "Não informada")
    return (city, dist, region, "Portugal")


rows = []
for _, r in d.iterrows():
    city, dist, region, country = parse_location(r["location"])
    size_label, size_order = size_map[r["size"]]
    loc_clean = ", ".join([x for x in [city, dist, country] if not x.startswith(("Não", "N/A"))])
    founded = int(r["founded"]) if r["founded"] else None
    rows.append([
        r["name"], r["url"] or None, r["industry"], r["social page"],
        loc_clean, city, dist, region, country, founded, size_label, size_order,
    ])

headers = ["Nome", "URL", "Indústria (Segmento)", "Página da Rede Social", "Localização",
           "Cidade", "Distrito", "Região", "País", "Fundação", "Tamanho", "Ordem do Tamanho"]
clean = pd.DataFrame(rows, columns=headers)
clean.to_csv(out_csv, index=False, encoding="utf-8-sig")

# ---------- Excel ----------
F = "Arial"
wb = Workbook()
ws = wb.active
ws.title = "Base_Limpa"
hdr_fill = PatternFill("solid", fgColor="1F3864")
thin = Side(style="thin", color="D9D9D9")

ws.append(headers)
for row in rows:
    ws.append(row)
for c in ws[1]:
    c.font = Font(name=F, bold=True, color="FFFFFF", size=10)
    c.fill = hdr_fill
    c.alignment = Alignment(vertical="center", wrap_text=True)
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name=F, size=10)
        c.border = Border(bottom=thin)
    row[9].number_format = "0"      # Fundação sem separador de milhar
    row[11].number_format = "0"
widths = [42, 40, 34, 42, 40, 24, 22, 28, 16, 11, 24, 12]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
ws.row_dimensions[1].height = 30

# ---------- Notas ----------
nt = wb.create_sheet("Notas_Limpeza")
last = len(rows) + 1
nt.column_dimensions["A"].width = 44
nt.column_dimensions["B"].width = 100


def title(r, text):
    nt.cell(r, 1, text).font = Font(name=F, bold=True, size=12, color="1F3864")


def line(r, a, b=None, bold=False):
    ca = nt.cell(r, 1, a)
    ca.font = Font(name=F, size=10, bold=bold)
    ca.alignment = Alignment(wrap_text=True, vertical="top")
    if b is not None:
        cb = nt.cell(r, 2, b)
        cb.font = Font(name=F, size=10)
        cb.alignment = Alignment(wrap_text=True, vertical="top")


r = 1
title(r, "Resumo da base (calculado a partir da aba Base_Limpa)"); r += 1
summary = [
    ("Total de empresas", f"=COUNTA(Base_Limpa!A2:A{last})"),
    ("Empresas com sede em Portugal", f'=COUNTIF(Base_Limpa!I2:I{last},"Portugal")'),
    ("Empresas fora de Portugal", f'=COUNTA(Base_Limpa!A2:A{last})-COUNTIF(Base_Limpa!I2:I{last},"Portugal")'),
    ("Sem cidade/região informada (só consta 'Portugal')", f'=COUNTIF(Base_Limpa!H2:H{last},"Não informada")'),
    ("Sem ano de fundação", f"=COUNTBLANK(Base_Limpa!J2:J{last})"),
    ("Sem URL do site", f"=COUNTBLANK(Base_Limpa!B2:B{last})"),
]
for a, f in summary:
    line(r, a, f)
    nt.cell(r, 2).alignment = Alignment(horizontal="left")
    r += 1

r += 1
title(r, "O que foi corrigido"); r += 1
fixes = [
    ("Coluna Tamanho (erro principal)",
     "Ao abrir a base, o Excel/Sheets leu faixas como datas: '11-50' virou 'nov/50' (95 empresas) e '1-10' virou '01/out.' (26 empresas). "
     "Foram restauradas e reescritas como texto que não vira data: '11 a 50' e '1 a 10'. As demais faixas foram escritas no mesmo padrão."),
    ("Ordem do Tamanho", "Coluna numérica (0 a 8) para o painel ordenar as faixas do menor para o maior. Sem ela, o Looker ordena em ordem alfabética ('1.001' antes de '11')."),
    ("Localização separada em colunas", "O texto original ('Cidade. Distrito. País') foi dividido em Cidade, Distrito e País. A coluna Localização ficou no formato 'Cidade, Distrito, País'."),
    ("Lisbon / Lisboa", "Padronizado para 'Lisboa' (aparecia dos dois jeitos, o que dividiria a mesma cidade em duas no painel)."),
    ("Coluna Região (nova)", "Criada a partir do Distrito para permitir o filtro por região: Norte, Centro, Área Metropolitana de Lisboa, Algarve e Açores."),
    ("Espaços extras", "Removidos espaços no início, no fim e duplicados em todos os campos."),
    ("Duplicados", "Verificado por nome, URL e página da rede social: nenhum duplicado encontrado. Nenhuma linha foi removida (388 empresas, como no original)."),
]
for a, b in fixes:
    line(r, a, b); r += 1

r += 1
title(r, "Decisões e suposições (confira se concorda)"); r += 1
assumptions = [
    ("Santarém → Centro", "Região classificada como 'Centro' (2 empresas). O distrito é dividido entre Centro e Alentejo, então é uma escolha."),
    ("Setúbal → Área Metropolitana de Lisboa", "Agrupado com Lisboa (6 empresas)."),
    ("10 empresas só com 'Portugal'", "Sem cidade nem distrito no original. Ficaram como 'Não informada', não foram inventadas."),
    ("'Greater Póvoa de Varzim Area' (2) e 'Porto Metropolitan Area' (1)", "Convertidas para Póvoa de Varzim / Porto, distrito Porto, região Norte."),
    ("'Vila do Porto' (Ilha de Santa Maria)", "Classificada em Açores (2 empresas; uma delas só trazia 'Vila do Porto. Portugal')."),
    ("3 empresas fora de Portugal", "Mantidas na base (EUA: 2, Brasil: 1), marcadas como 'Fora de Portugal'. Para mapear só empresas em Portugal, filtre País = Portugal."),
    ("Autônomo", "A categoria 'Self' (1 empresa) ficou como 'Autônomo (self-employed)'."),
    ("Indústria", "Mantida no original em inglês (94 segmentos), sem tradução."),
]
for a, b in assumptions:
    line(r, a, b); r += 1

r += 1
title(r, "Pontos de atenção na base original (não alterados)"); r += 1
attention = [
    ("'Póvoa de Lisboa' no distrito de Viseu", "21 empresas aparecem com essa cidade e distrito Viseu. Não consegui confirmar que esse local existe com esse nome; pode ser um erro da fonte. Mantido como veio, vale mencionar no relatório."),
    ("Ano de fundação em branco", "29 empresas sem ano. Ficaram em branco (não foi colocado zero)."),
    ("Anos muito antigos", "Há fundações entre 1615 e 1950 (24 empresas, a mais antiga de 1615). Parecem plausíveis para vinícolas, bancos e associações, e foram mantidos."),
    ("1 empresa sem URL", "A empresa sem URL também está sem ano."),
]
for a, b in attention:
    line(r, a, b); r += 1

r += 1
title(r, "Tamanho: de/para"); r += 1
line(r, "Valor que veio na base", "Valor na Base_Limpa (Ordem)", bold=True); r += 1
for k, (v, o) in sorted(size_map.items(), key=lambda kv: kv[1][1]):
    line(r, k, f"{v}  ({o})"); r += 1

wb.save(out_xlsx)
print("ok", len(rows))
print(clean["Região"].value_counts().to_dict())
print(clean["Tamanho"].value_counts().to_dict())
print(clean["Distrito"].value_counts().to_dict())
print("sem ano", clean["Fundação"].isna().sum(), "antes de 1950", (clean["Fundação"] < 1950).sum())
print("Viseu", (clean["Distrito"] == "Viseu").sum(), "Açores", (clean["Distrito"] == "Açores").sum())
print("Setubal", (clean["Distrito"] == "Setúbal").sum(), "Santarem", (clean["Distrito"] == "Santarém").sum())
